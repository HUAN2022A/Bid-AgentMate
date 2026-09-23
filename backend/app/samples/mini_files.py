"""样例迷你文件：python-docx 运行时生成三份小文件（招标/规范书/投标），落 storage 共享。

与 sample_tender_analysis() 的内容互为镜像；mini-bid 用真实 Heading 样式，
投标解析走标题切树路径（零 LLM）——样例模式全程无模型即可演示导入闭环。
错别字（"按装调试"）是刻意埋点，供 M3 错别字检查演示。
"""
import io

from docx import Document

from app.core.storage import storage

SAMPLE_FILES: list[tuple[str, str, str]] = [
    # (storage 相对路径, 上传文件名, role)
    ("samples/mini-tender.docx", "样例-翻车机机器人采购招标文件.docx", "main"),
    ("samples/mini-spec.docx", "样例-翻车机技术规范书.docx", "spec"),
    ("samples/mini-bid.docx", "样例-投标文件.docx", "bid"),
]


def _p(doc: Document, text: str, heading: int | None = None):
    if heading:
        doc.add_heading(text, level=heading)
    else:
        doc.add_paragraph(text)


def _mini_tender() -> bytes:
    doc = Document()
    _p(doc, "淮北国安电厂二期项目2×660MW超超临界机组翻车机自动摘钩、正钩、复钩机器人系统采购 招标文件（样例）", heading=1)
    _p(doc, "招标编号：JAHB-2026-SAMPLE")
    _p(doc, "第一章 投标邀请", heading=1)
    _p(doc, "本项目采购翻车机自动摘钩、正钩、复钩机器人系统 1 套，交货期：合同签订后 120 天。本项目不接受联合体投标。")
    _p(doc, "第二章 投标人须知", heading=1)
    _p(doc, "资格要求：投标人须具有近三年（2023 年起）至少 1 个翻车机或同类铁路装卸系统业绩。")
    _p(doc, "付款方式：合同签订后预付 30%，到货验收后付 60%，质保期满付 10%。")
    _p(doc, "废标条款：", heading=2)
    _p(doc, "D1 未按规定格式填写投标函或投标函未加盖法人公章的，其投标将被否决。")
    _p(doc, "D2 投标文件载明的招标项目完成期限超过招标文件规定期限的，其投标将被否决。")
    _p(doc, "D3 技术卷中出现报价信息的，其投标将被否决。")
    _p(doc, "第三章 评标办法", heading=1)
    table = doc.add_table(rows=6, cols=4)
    table.style = "Table Grid"
    rows = [
        ("编号", "评分项", "分值", "评分标准"),
        ("S1", "总体方案与技术路线", "10", "技术方案总体思路清晰，系统架构合理，与招标需求匹配，得 8-10 分；基本满足得 4-7 分。"),
        ("S2", "自动摘钩机器人技术性能", "15", "摘钩机器人满足 ★作业周期≤45秒、定位精度≤±5mm，全部满足得 12-15 分；一项不满足得 0 分。"),
        ("S3", "实施方案与进度保障", "8", "施工组织方案完善、进度计划合理、保障措施充分，得 6-8 分。"),
        ("S4", "质保期与售后服务", "5", "质保期满足 24 个月，售后响应≤4小时，得 4-5 分。"),
        ("S5", "投标报价", "30", "以有效投标中的最低评标价为基准价，基准价得满分，每高 1% 扣 0.5 分。"),
    ]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    _p(doc, "技术文件应包含：总体方案、设备技术性能、实施方案、质保与服务。技术文件一式 4 份，正本 1 份副本 3 份，A4 双面打印。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _mini_spec() -> bytes:
    doc = Document()
    _p(doc, "翻车机自动摘正复钩机器人系统技术规范书（样例）", heading=1)
    _p(doc, "2.1 ★机器人单钩摘钩作业周期 ≤45 秒（含定位、摘钩、复位）。", heading=2)
    _p(doc, "2.2 ★摘钩定位精度 ≤±5mm，重复定位精度 ≤±3mm。")
    _p(doc, "3.1 设备工作环境温度 -25℃~50℃，防护等级 IP65。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _mini_bid() -> bytes:
    doc = Document()
    _p(doc, "投标文件（商务技术文件）", heading=1)
    _p(doc, "一、投标函（不含报价）", heading=1)
    _p(doc, "致：淮北国安电厂：我方已仔细研究了贵方翻车机自动摘钩机器人系统采购招标文件的全部内容，愿以响应招标文件全部要求的方式投标。")
    _p(doc, "二、授权委托书", heading=1)
    _p(doc, "本授权书声明：我方法定代表人授权本项目项目经理为我方代理人，以我方名义处理本项目投标的一切事宜。")
    _p(doc, "八、投标设备技术性能指标的详细描述", heading=1)
    _p(doc, "8.1 总体方案", heading=2)
    _p(doc, "本方案面向翻车机自动摘钩作业场景，采用视觉定位与力控结合的技术路线：激光雷达扫描车皮钩头位置，六轴机械臂末端执行器完成自动摘钩，定位精度±5mm，满足招标要求。系统总体架构分为感知层、决策层与执行层。")
    _p(doc, "8.2 自动摘钩机器人技术性能", heading=2)
    _p(doc, "机器人单钩摘钩作业周期实测 42 秒（含定位、摘钩、复位），优于招标要求的 45 秒；重复定位精度±2mm，优于招标要求的±3mm。")
    _p(doc, "8.3 实施方案", heading=2)
    _p(doc, "项目实施分为设计、制造、按装调试、试运行四个阶段，总工期 110 天。现场安装调试配置 6 名技术人员，含 1 名项目经理。")  # 「按装」为错别字埋点
    _p(doc, "8.4 质保与服务", heading=2)
    _p(doc, "整机质保期 24 个月，质保期内 4 小时响应、24 小时到现场。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def ensure_sample_files() -> None:
    """三份迷你文件落 storage（幂等，首次生成后共享）。"""
    makers = {
        "samples/mini-tender.docx": _mini_tender,
        "samples/mini-spec.docx": _mini_spec,
        "samples/mini-bid.docx": _mini_bid,
    }
    for rel, maker in makers.items():
        if not storage.exists(rel):
            storage.put(rel, maker())
